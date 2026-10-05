"""
services/calendar_extract.py - DOCX / PDF / scanned-image parsers for the calendar import engine.

Every parser returns RAW ROWS (same shape as parse_csv) that go through the shared
normalize() -> validate() -> preview -> commit pipeline. Nothing here touches the database.
Extraction is a best guess: each row carries `confidence` and `raw_text` so the admin
can see what the row was read from.
"""
import io
import os
import re
from datetime import date

from services.calendar_import import (_ALIAS_LOOKUP, _norm_header, _rows_from_table,
                                      parse_date)

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
MONTHS.update({"sept": 9, "january": 1, "february": 2, "march": 3, "april": 4, "june": 6,
               "july": 7, "august": 8, "september": 9, "october": 10, "november": 11,
               "december": 12})
CONNECTOR = r"(?:to|till|until|through|-|–|—)"
SUFFIX = r"(?:st|nd|rd|th)?"
NUM_DATE = re.compile(r"(?<!\d)(\d{1,2})\s*[-/.]\s*(\d{1,2})\s*[-/.]\s*(\d{4}|\d{2})(?!\d)")
TXT_DATE = re.compile(r"(?<!\d)(\d{1,2})" + SUFFIX + r"[\s\-]*(?:of\s+)?([A-Za-z]{3,9})\.?(?:,?[\s\-]*(\d{4}))?(?!\d)")
SAME_MONTH_RANGE = re.compile(
    r"(?<!\d)(\d{1,2})" + SUFFIX + r"\s*" + CONNECTOR + r"\s*(\d{1,2})" + SUFFIX
    + r"[\s\-]*(?:of\s+)?([A-Za-z]{3,9})\.?,?\s*(\d{4})?(?!\d)", re.I)
OCR_NUM = re.compile(r"(?<![A-Za-z0-9])([0-9OoIl]{1,2})\s*([-/.])\s*([0-9OoIl]{1,2})\s*([-/.])\s*([0-9OoIl]{2,4})(?![A-Za-z0-9])")
WEEKDAY_PAREN = re.compile(r"\(\s*(?:mon|tue|wed|thu|fri|sat|sun)[a-z]*\s*\)", re.I)
FOLLOW = re.compile(r"(monday|tuesday|wednesday|thursday|friday|saturday)\s*(?:time\s*table|timetable|schedule)", re.I)

# (regex, event type code) - first match wins; order matters
KEYWORDS = [
    (r"last\s+working\s+day", "LAST_WORKING_DAY"),
    (r"(working|class(es)?)\s+(on\s+)?(saturday|sunday)|special\s+working|(saturday|sunday)\s+(is\s+)?(a\s+)?working", "SPECIAL_WORKING_DAY"),
    (r"attendance.*(freeze|closing|last date)", "ATTENDANCE_FREEZE"),
    (r"cie.*freeze|freeze", "CIE_FREEZE"),
    (r"marks?.*(submission|entry|upload|deadline)|submission of marks", "MARKS_DEADLINE"),
    (r"semester\s+end|see\b|end\s+sem|theory\s+exam|examination", "SEMESTER_EXAM"),
    (r"practical|lab\s+exam", "PRACTICAL_EXAM"),
    (r"internal\s+assessment|\bcie\b|\bia\s*[-–]?\s*\d|\bia\b|test\s*[-–]?\s*\d|mid[\s-]?term", "INTERNAL_ASSESSMENT"),
    (r"project\s+(review|phase|presentation)|mini\s+project", "PROJECT_REVIEW"),
    (r"parent|\bptm\b", "PTM"),
    (r"add\s*/?\s*drop", "ADD_DROP"),
    (r"withdraw", "WITHDRAWAL"),
    (r"registration", "REGISTRATION"),
    (r"feedback", "FACULTY_FEEDBACK"),
    (r"commencement|reopening|re-opening|first\s+day|begins?\b|start\s+of\s+(the\s+)?sem", "SEMESTER_START"),
    (r"vacation|recess|semester\s+break|mid[\s-]?sem\s+break|\bbreak\b", "NON_TEACHING_DAY"),
    (r"\bcca\b|co-?curricular", "CCA"),
    (r"holiday|jayanti|ugadi|diwali|deepavali|dussehra|dasara|ganesh|id[-\s]?(ul|e)|eid|christmas|good\s+friday|"
     r"republic\s+day|independence\s+day|gandhi|ambedkar|rajyotsava|kannada|sankranti|holi|ramzan|bakrid|"
     r"muharram|mahavir|buddha|basava|maha\s*shivaratri|ram\s*navami|may\s+day|labour\s+day|festival", "HOLIDAY"),
    (r"fest|cultural|sports|annual\s+day|college\s+day|convocation|workshop|seminar|orientation|induction", "COLLEGE_EVENT"),
]
KEYWORDS = [(re.compile(p, re.I), c) for p, c in KEYWORDS]


# ----------------------------------------------------------- date extraction
def _fix_ocr_digits(text):
    """O/o->0, I/l->1 inside numeric dates only. Returns (text, changed)."""
    changed = [False]

    def fix(m):
        parts = [m.group(1), m.group(3), m.group(5)]
        new = [p.replace("O", "0").replace("o", "0").replace("I", "1").replace("l", "1") for p in parts]
        if new != parts:
            changed[0] = True
        return f"{new[0]}{m.group(2)}{new[1]}{m.group(4)}{new[2]}"
    return OCR_NUM.sub(fix, text), changed[0]


def _year_for(day, month, year, hint):
    """Pick a year for a date written without one, using the academic period."""
    if year:
        return (2000 + year if year < 100 else year), 1.0
    if not hint:
        return None, 0.0
    ps, pe = hint
    inside = []
    for y in sorted({ps.year, pe.year}):
        try:
            d = date(y, month, day)
        except ValueError:
            continue
        if ps <= d <= pe:
            inside.append(y)
    if inside:
        return inside[0], 0.95
    return ps.year, 0.5


def _mk(day, month, year, hint):
    y, c = _year_for(day, month, year, hint)
    if y is None:
        return None, 0.0
    try:
        return date(y, month, day), c
    except ValueError:
        return None, 0.0


def find_dates(text, hint=None):
    """Returns list of (start, end, confidence, (span_start, span_end))."""
    found, taken = [], []

    def free(a, b):
        return all(b <= x or a >= y for x, y in taken)

    for m in SAME_MONTH_RANGE.finditer(text):
        mon = MONTHS.get(m.group(3).lower())
        if not mon or not free(*m.span()):
            continue
        y = int(m.group(4)) if m.group(4) else None
        d1, c1 = _mk(int(m.group(1)), mon, y, hint)
        d2, c2 = _mk(int(m.group(2)), mon, y, hint)
        if d1 and d2 and d1 <= d2:
            found.append((d1, d2, min(c1, c2), m.span()))
            taken.append(m.span())

    toks = []
    for m in NUM_DATE.finditer(text):
        if not free(*m.span()):
            continue
        y = int(m.group(3))
        d, c = _mk(int(m.group(1)), int(m.group(2)), y, hint)
        if d:
            toks.append((m.start(), m.end(), d, c))
        else:
            toks.append((m.start(), m.end(), None, 0.0))
    for m in TXT_DATE.finditer(text):
        mon = MONTHS.get(m.group(2).lower())
        if not mon or not free(*m.span()) or any(not (m.end() <= a or m.start() >= b) for a, b, _, _ in toks):
            continue
        y = int(m.group(3)) if m.group(3) else None
        d, c = _mk(int(m.group(1)), mon, y, hint)
        toks.append((m.start(), m.end(), d, c) if d else (m.start(), m.end(), None, 0.0))
    toks.sort()

    i = 0
    while i < len(toks):
        a, b, d, c = toks[i]
        if d is None:
            i += 1
            continue
        if i + 1 < len(toks) and toks[i + 1][2]:
            between = text[b:toks[i + 1][0]]
            if re.fullmatch(r"\s*" + CONNECTOR + r"\s*", between, re.I) and d <= toks[i + 1][2]:
                found.append((d, toks[i + 1][2], min(c, toks[i + 1][3]), (a, toks[i + 1][1])))
                i += 2
                continue
        found.append((d, d, c, (a, b)))
        i += 1
    return sorted(found, key=lambda f: f[3][0])


def classify(text):
    """(event_type_code, confidence, follow_weekday)"""
    follow = None
    f = FOLLOW.search(text)
    if f:
        follow = f.group(1).title()
    for rx, code in KEYWORDS:
        if rx.search(text):
            return code, 0.9, follow
    return "ACADEMIC_EVENT", 0.6, follow


def _clean_title(text, spans):
    for a, b in sorted(spans, reverse=True):
        text = text[:a] + " " + text[b:]
    text = WEEKDAY_PAREN.sub(" ", text)
    text = re.sub(r"^[\s\d\.\)\-–—:•*|]+(?=[A-Za-z])", "", text.strip())       # bullets / numbering
    text = re.sub(r"\s*[|]+\s*", " - ", text)
    text = re.sub(r"\s{2,}", " ", text).strip(" -–—:,;|")
    return text[:150]


def rows_from_line(text, hint, conf_cap=1.0, counter=None):
    """One text line -> 0..n raw rows."""
    text = re.sub(r"\s+", " ", text or "").strip()
    if len(text) < 6:
        return []
    text, fixed = _fix_ocr_digits(text)
    dates = find_dates(text, hint)
    if not dates:
        return []
    title = _clean_title(text, [d[3] for d in dates])
    code, tconf, follow = classify(title or text)
    multi = len(dates) > 1 and all(d[0] == d[1] for d in dates)
    out = []
    for s, e, dconf, _ in dates:
        conf = min(conf_cap, dconf, tconf)
        if multi:
            conf = min(conf, 0.7)            # several loose dates in one line: ambiguous
        if fixed:
            conf = min(conf, 0.75)           # we corrected OCR digits ourselves
        out.append({"event_type": code, "title": title, "start_date": s, "end_date": e,
                    "follow_weekday": follow if code == "SPECIAL_WORKING_DAY" else None,
                    "confidence": round(conf, 2), "raw_text": text[:300]})
    return out


# ----------------------------------------------------------- tables / shared
def rows_from_table(table, hint, conf_cap=1.0):
    table = [[("" if c is None else str(c).replace("\n", " ").strip()) for c in r] for r in table or []]
    table = [r for r in table if any(r)]
    if not table:
        return []
    headers = [_ALIAS_LOOKUP.get(_norm_header(h)) for h in table[0]]
    out = []
    if sum(1 for h in headers if h) >= 2:                      # recognised header row
        body = table[1:]
        for r in body:
            cells = {h: v for h, v in zip(headers, r) if h and v}
            start_ok = parse_date(cells.get("start_date"))[0] is not None
            if cells.get("event_type") and start_ok:
                cells.update(confidence=conf_cap, raw_text=" | ".join(r)[:300])
                out.append(cells)
            else:
                text = " ".join([cells.get("start_date", ""), cells.get("end_date", ""),
                                 cells.get("title", ""), cells.get("event_type", "")] if cells
                                else r)
                out += rows_from_line(text, hint, conf_cap)
        return out
    for r in table:                                            # headerless: treat each row as a line
        out += rows_from_line(" ".join(c for c in r if c), hint, conf_cap)
    return out


def _finish(rows):
    seen, out = set(), []
    for r in rows:
        s, e = parse_date(r.get("start_date"))[0], parse_date(r.get("end_date") or r.get("start_date"))[0]
        key = (str(s), str(e), str(r.get("title", "")).lower())
        if key in seen:
            continue
        seen.add(key)
        r["line"] = len(out) + 1
        out.append(r)
    if not out:
        raise ValueError("No dated entries were found in this file. Use the CSV template, or check "
                         "that the calendar lists dates as text (not only as colours in a grid).")
    return out


# --------------------------------------------------------------------- DOCX
def parse_docx(data, period=None):
    try:
        import docx
    except ImportError:
        raise ValueError("DOCX import needs python-docx:  pip install python-docx")
    hint = (period["start_date"], period["end_date"]) if period else None
    document = docx.Document(io.BytesIO(data))
    rows = []
    for t in document.tables:
        rows += rows_from_table([[c.text for c in r.cells] for r in t.rows], hint)
    covered = {(str(parse_date(r.get("start_date"))[0]), str(parse_date(r.get("end_date") or r.get("start_date"))[0]))
               for r in rows}
    for p in document.paragraphs:
        for r in rows_from_line(p.text, hint):
            if (str(r["start_date"]), str(r["end_date"])) not in covered:
                rows.append(r)
    return _finish(rows)


# ---------------------------------------------------------------------- OCR
def _ocr_lines(img):
    """PIL image -> [(line_text, confidence 0..1)] using Tesseract."""
    try:
        import pytesseract
    except ImportError:
        raise ValueError("Scanned/image import needs pytesseract:  pip install pytesseract "
                         "and the Tesseract program (set TESSERACT_CMD if it is not on PATH).")
    cmd = os.environ.get("TESSERACT_CMD")
    if cmd:
        pytesseract.pytesseract.tesseract_cmd = cmd
    try:
        d = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT, config="--psm 6")
    except pytesseract.TesseractNotFoundError:
        raise ValueError("Tesseract OCR is not installed or not found. Install it and/or set TESSERACT_CMD.")
    lines = {}
    for i, w in enumerate(d["text"]):
        if not str(w).strip():
            continue
        key = (d["block_num"][i], d["par_num"][i], d["line_num"][i])
        conf = float(d["conf"][i])
        lines.setdefault(key, []).append((w, conf if conf >= 0 else 0.0))
    out = []
    for key in sorted(lines):
        words = lines[key]
        out.append((" ".join(w for w, _ in words), sum(c for _, c in words) / len(words) / 100.0))
    return out


def parse_image(data, period=None):
    from PIL import Image
    hint = (period["start_date"], period["end_date"]) if period else None
    rows = []
    for text, conf in _ocr_lines(Image.open(io.BytesIO(data)).convert("L")):
        rows += rows_from_line(text, hint, conf_cap=min(conf, 0.95))
    return _finish(rows)


# ---------------------------------------------------------------------- PDF
def parse_pdf(data, period=None, max_pages=30):
    try:
        import pdfplumber
    except ImportError:
        raise ValueError("PDF import needs pdfplumber:  pip install pdfplumber")
    hint = (period["start_date"], period["end_date"]) if period else None
    rows, scanned_pages = [], []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for n, page in enumerate(pdf.pages[:max_pages], 1):
            tables = page.extract_tables() or []
            page_rows = []
            for t in tables:
                page_rows += rows_from_table(t, hint)
            text = (page.extract_text() or "").strip()
            if len(text) < 30 and not page_rows:
                scanned_pages.append(n)                       # no text layer -> OCR below
                continue
            got = {(str(parse_date(r.get("start_date"))[0]), str(parse_date(r.get("end_date") or r.get("start_date"))[0]))
                   for r in page_rows}
            for line in text.splitlines():
                for r in rows_from_line(line, hint):
                    if (str(r["start_date"]), str(r["end_date"])) not in got:
                        page_rows.append(r)
            rows += page_rows
    if scanned_pages:
        try:
            import pypdfium2 as pdfium
        except ImportError:
            raise ValueError("Scanned PDF import needs pypdfium2:  pip install pypdfium2")
        doc = pdfium.PdfDocument(data)
        for n in scanned_pages:
            img = doc[n - 1].render(scale=300 / 72).to_pil().convert("L")
            for text, conf in _ocr_lines(img):
                rows += rows_from_line(text, hint, conf_cap=min(conf, 0.95))
    return _finish(rows)